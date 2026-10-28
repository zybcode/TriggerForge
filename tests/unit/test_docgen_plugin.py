# TriggerForge - Event-driven directory orchestration engine.
# Copyright (C) 2026  [zybcode]
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version. See <https://www.gnu.org/licenses/>.

"""docgen 插件（文档自动生成）的子进程级测试。

覆盖部署方案 V1.0-R-03 的关键验收点：
- 正常生成：{{key}} 渲染、多条目合并 sections==1、成功统计产物
- 三态判定：模板缺失 → partial；渲染语法错误 → failed（V-4）
- V-3：未配置的 {{key}} 保留原文并列出
- 工程名清洗：Windows 保留名兜底
"""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import pytest

PLUGIN = Path(__file__).resolve().parents[2] / "plugins" / "docgen" / "main.py"


def _docxtpl_available() -> bool:
    try:
        import docxtpl  # noqa: F401
    except ImportError:
        return False
    return True


pytestmark = pytest.mark.skipif(
    not _docxtpl_available(), reason="docxtpl/python-docx not installed"
)


def _build_template(path: Path, extra_line: str | None = None) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("工程名称：{{工程名称}}")
    doc.add_paragraph("部位：{{分部工程部位}} 时间：{{分部工程时间}}")
    if extra_line:
        doc.add_paragraph(extra_line)
    doc.save(str(path))


def _write_txt(path: Path, content: str) -> None:
    path.write_text(content, encoding="utf-8-sig")


def _run_plugin(txt: Path, params: dict) -> dict:
    proc = subprocess.run(
        [sys.executable, str(PLUGIN), str(txt), json.dumps(params, ensure_ascii=False)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        check=False,
        timeout=120,
    )
    assert proc.returncode == 0, f"plugin crashed: {proc.stderr[-2000:]}"
    lines = [line for line in proc.stdout.strip().splitlines() if line.strip()]
    assert lines, f"plugin produced no stdout: {proc.stderr[-2000:]}"
    return json.loads(lines[-1])


@pytest.fixture()
def docgen_env(tmp_path: Path) -> dict:
    tpl_dir = tmp_path / "templates" / "验收模板"
    tpl_dir.mkdir(parents=True)
    watch = tmp_path / "watch"
    watch.mkdir()
    out_root = tmp_path / "shared"
    out_root.mkdir()
    return {
        "templates_root": str(tmp_path / "templates"),
        "output_root": str(out_root),
        "watch": watch,
        "out_root": out_root,
        "tpl_dir": tpl_dir,
    }


def test_docgen_success_and_merge(docgen_env: dict):
    """正常生成：多段落多条目合并 sections==1，产物齐全。"""
    _build_template(docgen_env["tpl_dir"] / "0.开工报告.docx", "未配置项：{{未配置项}}")
    txt = docgen_env["watch"] / "验证工程A.txt"
    _write_txt(
        txt,
        "模板类型=永安农业验收模板\n模板目录=验收模板\n"
        "[全局替换]\n工程名称=验证工程A\n"
        "[开工报告]\n模板文件=0.开工报告.docx\n"
        "[分部工程]\n模板文件=0.开工报告.docx\n"
        "基础柱=2025年12月15日\n主体梁=2025年12月20日\n",
    )

    params = {
        "templates_root": docgen_env["templates_root"],
        "output_root": docgen_env["output_root"],
    }
    result = _run_plugin(txt, params)

    assert result["status"] == "success"
    assert result["detail"] == "success"
    assert result["project"] == "验证工程A"

    folder = Path(result["output_folder"])
    merged = folder / "0.开工报告_分部工程.docx"
    assert merged.exists()

    from docx import Document

    doc = Document(str(merged))
    assert len(doc.sections) == 1
    body_text = "\n".join(p.text for p in doc.paragraphs)
    assert "2025年12月15日" in body_text
    assert "2025年12月20日" in body_text

    stats = (folder / "成功统计.txt").read_text(encoding="utf-8-sig")
    assert "分部工程：2 项" in stats
    # V-3：模板中未配置的占位符保留原文并列出
    assert "{{未配置项}}" in stats
    # 原配置档归档副本恢复原名
    assert (folder / "验证工程A.txt").exists()


def test_docgen_partial_on_missing_template(docgen_env: dict):
    """模板缺失 → partial：双产物并存，失败摘要具体（R2）。"""
    _build_template(docgen_env["tpl_dir"] / "0.开工报告.docx")
    txt = docgen_env["watch"] / "验证工程B.txt"
    _write_txt(
        txt,
        "模板类型=永安农业验收模板\n模板目录=验收模板\n"
        "[全局替换]\n工程名称=验证工程B\n"
        "[开工报告]\n模板文件=0.开工报告.docx\n"
        "[缺失段]\n模板文件=不存在.docx\n条目一=参数\n",
    )

    params = {
        "templates_root": docgen_env["templates_root"],
        "output_root": docgen_env["output_root"],
    }
    result = _run_plugin(txt, params)

    assert result["status"] == "success"
    assert result["detail"] == "partial"
    assert result["partial"] is True

    folder = Path(result["output_folder"])
    assert (folder / "成功统计.txt").exists()
    assert (folder / "错误原因.txt").exists()
    err = (folder / "错误原因.txt").read_text(encoding="utf-8-sig")
    assert "部分段落失败" in err
    assert "缺失段：模板缺失" in err


def test_docgen_failed_on_render_syntax_error(docgen_env: dict):
    """{{key} 少右括号 → TemplateSyntaxError → failed（V-4 / 验证项 32）。"""
    from docx import Document

    bad = Document()
    bad.add_paragraph("资料问题函 {{工程名称}")
    bad.save(str(docgen_env["tpl_dir"] / "18.资料问题函.docx"))

    txt = docgen_env["watch"] / "验证工程C.txt"
    _write_txt(
        txt,
        "模板类型=永安农业验收模板\n模板目录=验收模板\n"
        "[全局替换]\n工程名称=验证工程C\n"
        "[资料问题函]\n模板文件=18.资料问题函.docx\n第一条=参数\n",
    )

    params = {
        "templates_root": docgen_env["templates_root"],
        "output_root": docgen_env["output_root"],
    }
    result = _run_plugin(txt, params)

    assert result["status"] == "failed"
    assert "TemplateSyntaxError" in result["error"]
    assert "unexpected '}'" in result["error"]

    folder = Path(result["output_folder"])
    assert (folder / "错误原因.txt").exists()
    err = (folder / "错误原因.txt").read_text(encoding="utf-8-sig")
    assert "【模板渲染异常明细】" in err
    assert "【模板占位符规范提示】" in err


def test_docgen_empty_config_is_failed(docgen_env: dict):
    """空配置/无段落 → failed（V-1），摘要明确。"""
    txt = docgen_env["watch"] / "空配置.txt"
    _write_txt(txt, "模板类型=永安农业验收模板\n模板目录=验收模板\n")

    params = {
        "templates_root": docgen_env["templates_root"],
        "output_root": docgen_env["output_root"],
    }
    result = _run_plugin(txt, params)

    assert result["status"] == "failed"
    assert "未识别到任何有效段落" in result["error"]


def test_docgen_sanitize_windows_reserved_name(docgen_env: dict):
    """COM1 等保留名兜底 _未识别_（T4 口径）。"""
    _build_template(docgen_env["tpl_dir"] / "0.开工报告.docx")
    txt = docgen_env["watch"] / "保留名工程.txt"
    _write_txt(
        txt,
        "模板类型=永安农业验收模板\n模板目录=验收模板\n"
        "[全局替换]\n工程名称=COM1\n"
        "[开工报告]\n模板文件=0.开工报告.docx\n",
    )

    params = {
        "templates_root": docgen_env["templates_root"],
        "output_root": docgen_env["output_root"],
    }
    result = _run_plugin(txt, params)

    assert result["status"] == "success"
    assert result["project"] == "_未识别_"


def test_docgen_singleton_template(docgen_env: dict):
    """段落无数据键 → 虚拟条目渲染 1 次（R4 单例模板）。"""
    _build_template(docgen_env["tpl_dir"] / "0.开工报告.docx")
    txt = docgen_env["watch"] / "单例工程.txt"
    _write_txt(
        txt,
        "模板类型=永安农业验收模板\n模板目录=验收模板\n"
        "[全局替换]\n工程名称=单例工程\n"
        "[开工报告]\n模板文件=0.开工报告.docx\n",
    )

    params = {
        "templates_root": docgen_env["templates_root"],
        "output_root": docgen_env["output_root"],
    }
    result = _run_plugin(txt, params)

    assert result["status"] == "success"
    folder = Path(result["output_folder"])
    assert (folder / "0.开工报告_开工报告.docx").exists()
    stats = (folder / "成功统计.txt").read_text(encoding="utf-8-sig")
    assert "开工报告：1 项" in stats
