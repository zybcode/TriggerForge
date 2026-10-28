# TriggerForge 插件：docgen —— 文档自动生成（Word 模板批量渲染合并）
# 实现《文档自动生成系统部署方案 V1.0-R-03》阐述的全部核心功能：
#   1. TXT 配置档解析（编码回退链 utf-8-sig → gbk → gb18030）
#   2. 工程名清洗（非法字符 / Windows 保留名兜底 _未识别_）
#   3. 模板目录解析：[root] 模板目录= 默认子目录，段内 模板目录= 可覆盖（R3/R7）
#   4. docxtpl {{key}} 占位符渲染（V-4 锁定口径：禁止 ${key} / {key} / {{key}）
#   5. 多条目合并输出（append_doc 保留 sectPr，合并后 sections == 1）
#   6. 单例模板：段落无数据键时构造虚拟条目渲染 1 次（R4）
#   7. 未替换占位符检测：正文/表格(含嵌套)/页眉/页脚 全覆盖（V-3 / R19）
#   8. 三态判定（V-1/V-2/V-4）：success / partial / failed
#   9. 三态产物：成功统计.txt / 错误原因.txt（原子写，utf-8-sig），区块与方案附录 B 一致
#
# TriggerForge 插件契约：
#   argv[1] = 投递的 TXT 配置文件路径
#   argv[2] = JSON 参数（亦可通过环境变量 TRIGGERFORGE_PARAMS 注入），支持：
#       templates_root     模板根目录（模板类型未登记时的兜底根）
#       global_config_path 全局配置 JSON 路径（模板类型 → {"模板目录": 根目录}）
#       output_root        工程输出根目录（默认：TXT 所在目录）
#       min_fee            未替换占位符检测的豁免键（默认 "min_fee"）
#   stdout 必须输出单个 JSON；status ∈ {success, failed}。
#   三态映射：success/partial → status=success（partial 以 partial=true 标注）；
#             failed → status=failed（TriggerForge 据此归档到 error_path）。

from __future__ import annotations

import json
import logging
import os
import re
import shutil
import sys
import time
import traceback

logger = logging.getLogger("docgen")

# 段落"失败"备注集合（三态判据 + 产物分区统一引用，V-4 / R20）
_FAIL_NOTES = ("模板缺失", "路径非法", "缺少模板文件", "模板渲染异常")

_WIN_RESERVED = re.compile(r"(?i)^(COM|LPT)[1-9]$")
_WIN_RESERVED_NAMES = {"CON", "PRN", "AUX", "NUL"}


# ----------------------------------------------------------------------
# 参数加载
# ----------------------------------------------------------------------
def _load_params(argv: list[str]) -> dict:
    """优先 argv[2] JSON，回退环境变量 TRIGGERFORGE_PARAMS。"""
    if len(argv) > 2 and argv[2].strip():
        try:
            parsed = json.loads(argv[2])
            if isinstance(parsed, dict):
                return parsed
        except json.JSONDecodeError:
            logger.warning("argv[2] 不是合法 JSON，尝试环境变量参数")
    env_raw = os.environ.get("TRIGGERFORGE_PARAMS", "")
    if env_raw.strip():
        try:
            parsed = json.loads(env_raw)
            if isinstance(parsed, dict):
                return parsed
        except json.JSONDecodeError:
            logger.warning("TRIGGERFORGE_PARAMS 不是合法 JSON，忽略")
    return {}


# ----------------------------------------------------------------------
# TXT 配置解析（编码回退链）
# ----------------------------------------------------------------------
def load_txt_config(txt_path: str) -> dict:
    """解析 INI 风格 TXT 配置；编码回退 utf-8-sig → gbk → gb18030。"""
    if not os.path.exists(txt_path):
        raise FileNotFoundError(f"配置文件不存在：{txt_path}")

    text = None
    used_enc = None
    for enc in ("utf-8-sig", "gbk", "gb18030"):
        try:
            with open(txt_path, "r", encoding=enc) as f:
                text = f.read()
            used_enc = enc
            break
        except UnicodeDecodeError:
            continue

    if text is None:
        raise UnicodeDecodeError(
            "TXT", b"", 0, 1, "无法识别编码（试过 utf-8-sig / gbk / gb18030）"
        )

    logger.info("配置文件编码识别为 %s：%s", used_enc, txt_path)

    config: dict = {"root": {}}
    current_section = "root"
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("[") and line.endswith("]"):
            current_section = line[1:-1].strip()
            config.setdefault(current_section, {})
            continue
        if "=" in line:
            key, value = line.split("=", 1)
            config.setdefault(current_section, {})[key.strip()] = value.strip()
    return config


# ----------------------------------------------------------------------
# 工具函数
# ----------------------------------------------------------------------
def _sanitize_name(name: str) -> str:
    """剔除非法字符、拒绝 ..、精确匹配 Windows 保留名。"""
    for c in '<>:"/\\|?*':
        name = name.replace(c, "_")
    name = name.strip().strip(".")
    if (
        name in ("", ".", "..")
        or name.upper() in _WIN_RESERVED_NAMES
        or _WIN_RESERVED.match(name)
    ):
        name = "_未识别_"
    return name


def _safe_remove(path: str) -> None:
    try:
        if os.path.exists(path):
            os.remove(path)
    except OSError as e:
        logger.warning("删除失败 %s: %s", path, e)


def _atomic_write(path: str, text: str) -> None:
    """原子写：tmp + os.replace；失败清理 tmp 并抛出。"""
    tmp = path + ".tmp"
    try:
        with open(tmp, "w", encoding="utf-8-sig") as f:
            f.write(text)
        os.replace(tmp, path)
    except Exception:
        logger.exception("原子写入失败 %s", path)
        if os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass
        raise


def _archive_conflict_free(src: str, dst: str) -> str:
    """归档文件；同名冲突加时间戳后缀，返回最终路径。"""
    if os.path.exists(dst):
        ts = time.strftime("%Y%m%d%H%M%S")
        dst = f"{dst}.{ts}"
    shutil.copy2(src, dst)
    return dst


# ----------------------------------------------------------------------
# Word 渲染与合并
# ----------------------------------------------------------------------
def replace_placeholders(template_path: str, mapping: dict):
    """docxtpl 渲染：默认 Jinja2 分隔符（{{key}}，V-4）。返回底层 Document。

    未定义变量使用 _KeepPlaceholderUndefined 保留 {{key}} 原文，
    使输出侧 check_unreplaced_placeholders 能检出（V-3 锁定口径）。
    """
    from docxtpl import DocxTemplate
    from jinja2 import Environment, Undefined

    class _KeepPlaceholderUndefined(Undefined):
        """未定义变量渲染时保留 {{key}} 原文，便于残留检测。"""

        def __str__(self):
            return "{{" + str(self._undefined_name) + "}}"

    if not isinstance(template_path, str) or not os.path.exists(template_path):
        raise FileNotFoundError(f"找不到指定的 Word 模板文件路径: {template_path}")

    # 清理 mapping 键（兼容历史 ${...} / &{...} / {...} 包装写法）
    clean_mapping: dict = {}
    for k, v in mapping.items():
        if v is None:
            continue
        clean_key = str(k).strip()
        if clean_key.startswith("${") and clean_key.endswith("}"):
            clean_key = clean_key[2:-1].strip()
        elif clean_key.startswith("&{") and clean_key.endswith("}"):
            clean_key = clean_key[2:-1].strip()
        elif clean_key.startswith("{") and clean_key.endswith("}"):
            clean_key = clean_key[1:-1].strip()
        clean_mapping[clean_key] = v

    doc = DocxTemplate(template_path)
    env = Environment(undefined=_KeepPlaceholderUndefined)
    doc.render(clean_mapping, jinja_env=env)
    return doc.docx


def append_doc(main_doc, source_doc):
    """分页符 + 追加 source body 元素（过滤 body 级 sectPr）。

    按方案验收标准/T2 沙盒口径：合并后 len(doc.sections) == 1。
    main_doc 自身的 body 末尾 sectPr 保持不动（唯一节），source body
    级 sectPr 被过滤，条目间用分页符分隔。
    """
    import copy

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p = main_doc.add_paragraph()
    r = p.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    r._r.append(br)

    main_body = main_doc.element.body
    for el in list(source_doc.element.body):
        if el.tag == qn("w:sectPr"):
            continue
        main_body.append(copy.deepcopy(el))


def check_unreplaced_placeholders(file_path: str, exempt: str = "min_fee") -> list:
    """覆盖正文/表格(含嵌套)/页眉/页脚；返回未替换 key 列表（V-3 / R19）。"""
    from docx import Document

    try:
        doc = Document(file_path)
        pattern = re.compile(r"\{\{([^{}]+)\}\}")
        unreplaced: set = set()

        def _scan_paragraphs(paragraphs):
            for paragraph in paragraphs:
                for k in pattern.findall(paragraph.text):
                    if k and k != exempt:
                        unreplaced.add(k)

        def _scan_table(table):
            for row in table.rows:
                for cell in row.cells:
                    _scan_paragraphs(cell.paragraphs)
                    for nested in cell.tables:
                        _scan_table(nested)

        _scan_paragraphs(doc.paragraphs)
        for table in doc.tables:
            _scan_table(table)

        for section in doc.sections:
            for hf in (
                section.header,
                section.first_page_header,
                section.even_page_header,
                section.footer,
                section.first_page_footer,
                section.even_page_footer,
            ):
                if hf is None or hf.is_linked_to_previous:
                    continue
                _scan_paragraphs(hf.paragraphs)
                for t in hf.tables:
                    _scan_table(t)

        if unreplaced:
            final = ["{{" + p + "}}" for p in sorted(unreplaced)]
            logger.warning("⚠️ %s 残留占位符：%s", file_path, final)
        return sorted(unreplaced)
    except Exception as e:
        logger.error("检查未替换占位符异常：%s", e)
        return []


# ----------------------------------------------------------------------
# 段落生成（渲染异常入 stats，R18）
# ----------------------------------------------------------------------
def _summarize_render_errors(errors: list) -> str:
    """渲染异常摘要：首条异常类型+信息（共 N 条渲染失败）。"""
    if not errors:
        return ""
    first = errors[0]
    return (
        f"{first.get('异常类型')}: {first.get('异常信息')}"
        f"（共 {len(errors)} 条渲染失败）"
    )


def generate_single_file_multi_sections(
    template_path: str,
    template_root: str,
    output_folder: str,
    section_name: str,
    section_items: dict,
    global_map: dict,
    exempt_key: str = "min_fee",
) -> dict:
    import os as _os

    rel = _os.path.relpath(template_path, template_root)
    rel_dir = _os.path.dirname(rel)
    base_name = _os.path.splitext(_os.path.basename(rel))[0]
    output_name = f"{base_name}_{section_name}.docx"
    output_dir = _os.path.join(output_folder, rel_dir) if rel_dir else output_folder
    _os.makedirs(output_dir, exist_ok=True)
    output_path = _os.path.join(output_dir, output_name)

    main_doc = None
    data_count = 0
    unreplaced_by_file: dict = {}
    render_errors: list = []

    field_order = global_map.get("_field_order", ["时间", "部位", "位置", "备注"])

    # 单例模板：无数据键时构造虚拟条目，渲染一次（R4）
    if not section_items:
        section_items = {"": ""}

    for idx, (name, raw_value) in enumerate(section_items.items(), start=1):
        logger.info("[%s] 处理第 %d 项：%s → %s", section_name, idx, name, raw_value)

        mapping = dict(global_map)
        mapping["编号"] = f"{idx:03d}"
        mapping[f"{section_name}编号"] = f"{idx:03d}"
        mapping["序号"] = str(idx)
        mapping[section_name] = name

        if raw_value:
            parts = [p.strip() for p in str(raw_value).split("|")]
            for i, val in enumerate(parts):
                val = val.strip()
                if not val:
                    continue
                # 优先级 1：命名参数（中文标识符:值，如 规格:HRB400）
                m = re.match(r"^([\u4e00-\u9fa5][\u4e00-\u9fa5_]{0,9})[：:](.+)$", val)
                if m:
                    field_name = m.group(1)
                    field_value = m.group(2).strip()
                    mapping[f"{section_name}{field_name}"] = field_value
                    mapping[field_name] = field_value
                    continue
                # 优先级 2：位置参数（按 field_order）
                if i < len(field_order):
                    mapping[f"{section_name}{field_order[i]}"] = val
                    mapping[field_order[i]] = val
                else:
                    mapping[f"{section_name}参数{i}"] = val

        try:
            doc = replace_placeholders(template_path, mapping)
        except MemoryError:
            logger.exception("[%s] 第 %d 项渲染内存不足", section_name, idx)
            render_errors.append(
                {
                    "序号": idx,
                    "条目": str(name),
                    "异常类型": "MemoryError",
                    "异常信息": "内存不足",
                }
            )
            continue
        except Exception as e:
            logger.exception("[%s] 第 %d 项渲染出错", section_name, idx)
            render_errors.append(
                {
                    "序号": idx,
                    "条目": str(name),
                    "异常类型": type(e).__name__,
                    "异常信息": str(e),
                }
            )
            continue

        if doc is None:
            render_errors.append(
                {
                    "序号": idx,
                    "条目": str(name),
                    "异常类型": "NoneDoc",
                    "异常信息": "replace_placeholders 返回 None",
                }
            )
            continue

        try:
            if main_doc is None:
                main_doc = doc
            else:
                append_doc(main_doc, doc)
            data_count += 1
        except MemoryError:
            logger.exception("[%s] 第 %d 项合并内存不足", section_name, idx)
            render_errors.append(
                {
                    "序号": idx,
                    "条目": str(name),
                    "异常类型": "MemoryError",
                    "异常信息": "合并内存不足",
                }
            )
            break
        except Exception as e:
            logger.exception("[%s] 第 %d 项合并失败", section_name, idx)
            render_errors.append(
                {
                    "序号": idx,
                    "条目": str(name),
                    "异常类型": type(e).__name__,
                    "异常信息": str(e),
                }
            )
            continue

    if main_doc is not None:
        try:
            main_doc.save(output_path)
            logger.info("[%s] 文件已合并保存：%s", section_name, output_path)
            keys = check_unreplaced_placeholders(output_path, exempt_key)
            rel_out = _os.path.join(rel_dir, output_name) if rel_dir else output_name
            unreplaced_by_file[rel_out] = keys
        except MemoryError:
            logger.exception("[%s] 保存文件内存不足：%s", section_name, output_path)
            render_errors.append(
                {
                    "序号": "-",
                    "条目": "-",
                    "异常类型": "MemoryError",
                    "异常信息": "保存文件内存不足",
                }
            )
        except Exception as e:
            logger.exception("[%s] 保存文件失败：%s", section_name, output_path)
            render_errors.append(
                {
                    "序号": "-",
                    "条目": "-",
                    "异常类型": type(e).__name__,
                    "异常信息": f"保存失败：{e}",
                }
            )
    else:
        logger.warning("[%s] 未生成有效内容，跳过保存", section_name)

    seg_stats: dict = {
        "输出子目录": rel_dir,
        "数据条数": data_count,
        "文件": [output_name] if main_doc is not None else [],
        "未替换占位符": unreplaced_by_file,
    }
    if render_errors:
        seg_stats["渲染异常"] = render_errors
        if data_count == 0:
            seg_stats["备注"] = "模板渲染异常"
            seg_stats["异常摘要"] = _summarize_render_errors(render_errors)
    return seg_stats


def generate_pages_from_config(
    template_root: str,
    output_folder: str,
    config: dict,
    global_map: dict,
    exempt_key: str = "min_fee",
) -> dict:
    """多段落入口：[root] 模板目录 默认子目录 + 段内 模板目录= 覆盖（R3/R7）。"""
    stats: dict = {"段落统计": {}}
    skip_sections = {"root", "全局替换", "模板类型", "_说明"}

    default_sub = config.get("root", {}).get("模板目录", "").strip()
    if default_sub and (default_sub.startswith("..") or os.path.isabs(default_sub)):
        logger.error("[root] 模板目录非法：%s", default_sub)
        default_sub = ""

    for section_name, section_items in config.items():
        if section_name in skip_sections:
            continue

        template_file = section_items.get("模板文件")
        if not template_file:
            stats["段落统计"][section_name] = {
                "输出子目录": "",
                "数据条数": 0,
                "文件": [],
                "未替换占位符": {},
                "备注": "缺少模板文件",
            }
            logger.warning("[%s] 缺少 模板文件=xxx.docx，已跳过", section_name)
            continue

        seg_sub = section_items.get("模板目录", default_sub).strip()
        if seg_sub and (seg_sub.startswith("..") or os.path.isabs(seg_sub)):
            stats["段落统计"][section_name] = {
                "输出子目录": "",
                "数据条数": 0,
                "文件": [],
                "未替换占位符": {},
                "备注": "路径非法",
            }
            logger.error("[%s] 模板目录非法：%s", section_name, seg_sub)
            continue

        norm = os.path.normpath(template_file)
        if norm.startswith("..") or os.path.isabs(norm):
            stats["段落统计"][section_name] = {
                "输出子目录": "",
                "数据条数": 0,
                "文件": [],
                "未替换占位符": {},
                "备注": "路径非法",
            }
            logger.error("[%s] 模板文件路径非法：%s", section_name, template_file)
            continue

        seg_root = os.path.join(template_root, seg_sub) if seg_sub else template_root
        template_path = os.path.join(seg_root, norm)
        if not os.path.exists(template_path):
            stats["段落统计"][section_name] = {
                "输出子目录": os.path.dirname(norm),
                "数据条数": 0,
                "文件": [],
                "未替换占位符": {},
                "备注": "模板缺失",
            }
            logger.error("[%s] 模板文件不存在：%s", section_name, template_path)
            continue

        data = {
            k: v for k, v in section_items.items() if k not in ("模板文件", "模板目录")
        }
        stats["段落统计"][section_name] = generate_single_file_multi_sections(
            template_path=template_path,
            template_root=seg_root,
            output_folder=output_folder,
            section_name=section_name,
            section_items=data,
            global_map=global_map,
            exempt_key=exempt_key,
        )

    return stats


# ----------------------------------------------------------------------
# 三态判定（V-1 / V-2 / V-4 锁定口径）
# ----------------------------------------------------------------------
def _judge_status(stats: dict) -> str:
    段落统计 = stats.get("段落统计", {})
    if not 段落统计:
        return "failed"  # V-1：空配置/无段落
    fail = sum(1 for i in 段落统计.values() if i.get("备注") in _FAIL_NOTES)
    succ = sum(1 for i in 段落统计.values() if i.get("数据条数", 0) > 0)
    if succ == 0:
        return "failed"
    if fail > 0:
        return "partial"
    return "success"


# ----------------------------------------------------------------------
# 产物区块
# ----------------------------------------------------------------------
def _format_unreplaced_block(段落统计: dict) -> list:
    """【未被替换的占位符】区块，每个模板文件都列（含"无"）。"""
    lines = []
    any_file = False
    for seg, info in 段落统计.items():
        unrepl = info.get("未替换占位符", {})
        for fname, keys in unrepl.items():
            any_file = True
            if keys:
                rendered = "、".join("{{" + k + "}}" for k in keys)
                lines.append(f"  {fname}：{rendered}")
            else:
                lines.append(f"  {fname}：无")
    if not any_file:
        lines.append("  （无模板文件产出）")
    return lines


def write_success_stats(save_folder: str, stats: dict, orig_name: str, partial=False):
    path = os.path.join(save_folder, "成功统计.txt")
    lines = []
    lines.append("=" * 60)
    lines.append("处理成功统计" if not partial else "处理部分成功统计")
    lines.append("=" * 60)
    lines.append(f"配置文件：{orig_name}")
    lines.append(f"工程名称：{stats.get('工程名称', '')}")
    lines.append(f"模板类型：{stats.get('模板类型', '')}")
    lines.append(f"处理时间：{time.strftime('%Y-%m-%d %H:%M:%S')}")
    lines.append("=" * 60)
    lines.append("")
    if partial:
        lines.append("⚠️ 部分段落未生成（详见同目录 错误原因.txt）")
        lines.append("")

    段落统计 = stats.get("段落统计", {})
    lines.append("【成功项（按段落）】")
    total = 0
    for seg, info in 段落统计.items():
        n = info.get("数据条数", 0)
        if n > 0:
            total += n
            lines.append(f"{seg}：{n} 项")
    lines.append(f"合计：{total} 项")
    lines.append("")

    # 零数据段落（合法，无数据）—— 排除所有失败备注（V-2 / Q-背景-2）
    zero_segs = [
        (seg, info)
        for seg, info in 段落统计.items()
        if info.get("数据条数", 0) == 0 and info.get("备注") not in _FAIL_NOTES
    ]
    if zero_segs:
        lines.append("【零数据段落（合法，无数据）】")
        for seg, _ in zero_segs:
            lines.append(f"  {seg}：0 条")
        lines.append("")

    lines.append("【输出明细】")
    for seg, info in 段落统计.items():
        sub = info.get("输出子目录", "")
        files = info.get("文件", [])
        if not files:
            continue
        header = f"{sub}/" if sub else "（根级）"
        lines.append(header)
        for f in files:
            lines.append(f"  - {f}")
    lines.append("")

    # 渲染异常（部分条目未生成）—— 有数据但部分条目渲染失败（R22）
    render_error_segs = [
        (seg, info)
        for seg, info in 段落统计.items()
        if info.get("渲染异常") and info.get("数据条数", 0) > 0
    ]
    if render_error_segs:
        lines.append("【渲染异常（部分条目未生成）】")
        for seg, info in render_error_segs:
            for err in info.get("渲染异常", []):
                lines.append(
                    f"  {seg} 条目#{err['序号']} {err['条目']}："
                    f"{err['异常类型']}: {err['异常信息']}"
                )
        lines.append("")

    lines.append("【未被替换的占位符】")
    lines.extend(_format_unreplaced_block(段落统计))
    lines.append("")
    lines.append("=" * 60)
    lines.append("处理完成。请核对“未被替换的占位符”一项。")
    lines.append("=" * 60)

    _atomic_write(path, "\n".join(lines))
    logger.info("成功统计已写入：%s", path)


def write_error_reason(save_folder: str, stats: dict, orig_name: str, partial=False):
    path = os.path.join(save_folder, "错误原因.txt")
    lines = []
    lines.append("=" * 60)
    lines.append("处理失败报告" if not partial else "处理部分失败报告")
    lines.append("=" * 60)
    lines.append(f"配置文件：{orig_name}")
    lines.append(f"工程名称：{stats.get('工程名称', '')}")
    lines.append(f"模板类型：{stats.get('模板类型', '')}")
    lines.append(f"处理时间：{time.strftime('%Y-%m-%d %H:%M:%S')}")
    lines.append("=" * 60)
    lines.append("")
    lines.append("【失败摘要】")
    lines.append(stats.get("异常", "未知异常"))
    lines.append("")
    lines.append("【可能原因】")
    lines.append("1. 模板文件缺失或路径错误")
    lines.append("2. 配置档缺少对应字段")
    lines.append("3. 模板占位符语法错误（如 {{key} 少一个右花括号）或与配置不匹配")
    lines.append("4. 配置文件编码不受支持或内容为空")
    lines.append("")
    lines.append("【处理建议】")
    lines.append("1. 检查模板根目录下对应模板是否存在")
    lines.append("2. 核对配置档对应字段")
    lines.append("3. 确认 TXT 为 UTF-8（建议带 BOM）")
    lines.append("4. 若报 unexpected '}'，检查模板是否把 {{key}} 写成 {{key}")
    lines.append("5. 修正后重新投递本配置文件到监控目录")
    lines.append("")

    段落统计 = stats.get("段落统计", {})

    # 失败段落清单（含异常摘要，R23）
    lines.append("【失败段落】")
    any_fail = False
    for seg, info in 段落统计.items():
        note = info.get("备注", "")
        if note in _FAIL_NOTES:
            any_fail = True
            summary = info.get("异常摘要", "")
            line = f"  {seg}：{note}"
            if summary:
                line += f"（{summary}）"
            lines.append(line)
    if not any_fail:
        lines.append("  （无失败段落，或异常发生于段落解析之前）")
    lines.append("")

    # 模板渲染异常明细 + 占位符规范提示（R23 / V-4）
    render_error_segs = [
        (seg, info) for seg, info in 段落统计.items() if info.get("渲染异常")
    ]
    if render_error_segs:
        lines.append("【模板渲染异常明细】")
        for seg, info in render_error_segs:
            lines.append(f"  [{seg}]")
            for err in info.get("渲染异常", []):
                lines.append(
                    f"    条目#{err['序号']} {err['条目']}："
                    f"{err['异常类型']}: {err['异常信息']}"
                )
        lines.append("")
        lines.append("【模板占位符规范提示】")
        lines.append("1. 本系统模板占位符为 {{key}}（docxtpl 0.20.x 标准，前后各两个花括号）")
        lines.append("2. 正确：{{key}}")
        lines.append("3. 常见错误：")
        lines.append("   - {{key}   少一个右花括号 → 触发 unexpected '}'")
        lines.append("   - {key}    单花括号，不会被识别")
        lines.append("   - ${key}   旧语法，本系统已不识别（会原样输出）")
        lines.append("4. 修改模板后无需重启引擎，下一轮投递立即生效（模板热更新）")
        lines.append("")

    lines.append("【按模板文件的未替换占位符】")
    lines.extend(_format_unreplaced_block(段落统计))
    lines.append("")

    lines.append("【完整异常堆栈】")
    lines.append(stats.get("堆栈", "（无）"))
    lines.append("")
    lines.append("=" * 60)
    lines.append("如需协助，请将此文件发给管理员。")
    lines.append("=" * 60)

    _atomic_write(path, "\n".join(lines))
    logger.error("错误原因已写入：%s", path)


# ----------------------------------------------------------------------
# 单配置处理
# ----------------------------------------------------------------------
def _resolve_template_root(params: dict, template_type: str) -> str | None:
    """模板类型 → 模板根目录。global_config_path 优先，回退 templates_root。"""
    global_config_path = params.get("global_config_path")
    if global_config_path and template_type:
        try:
            with open(global_config_path, "r", encoding="utf-8") as f:
                data = json.load(f)
            entry = data.get(template_type)
            if isinstance(entry, dict) and entry.get("模板目录"):
                return entry["模板目录"]
        except Exception as e:
            logger.error("读取全局配置失败：%s", e)
            return None
    root = params.get("templates_root")
    return root if root else None


def process_config(txt_path: str, orig_name: str, params: dict) -> dict:
    """处理单个 TXT 配置档，返回 {detail, save_folder, stats}。"""
    output_root = params.get("output_root") or os.path.dirname(
        os.path.abspath(txt_path)
    )
    exempt_key = params.get("min_fee", "min_fee")

    # 1. 读取配置（异常边界单独捕获，尽力取工程名）
    try:
        config = load_txt_config(txt_path)
    except Exception as e:
        logger.exception("读取配置文件失败：%s", txt_path)
        project_name = _sanitize_name(os.path.splitext(orig_name)[0])
        save_folder = os.path.join(output_root, project_name or "_未识别_")
        os.makedirs(save_folder, exist_ok=True)
        stats = {
            "工程名称": project_name or "",
            "模板类型": "",
            "段落统计": {},
            "异常": f"读取配置文件失败：{e}",
            "堆栈": traceback.format_exc(),
        }
        write_error_reason(save_folder, stats, orig_name)
        return {"detail": "failed", "save_folder": save_folder, "stats": stats}

    # 2. 工程名解析 + 目录逃逸校验
    project_name = config.get("全局替换", {}).get("工程名称") or os.path.splitext(
        orig_name
    )[0]
    project_name = _sanitize_name(project_name)
    save_folder = os.path.abspath(os.path.join(output_root, project_name))
    output_root_abs = os.path.abspath(output_root)
    if not save_folder.startswith(output_root_abs):
        raise ValueError(
            f"工程目录 {save_folder} 超出输出根目录 {output_root_abs}"
        )
    os.makedirs(save_folder, exist_ok=True)

    # 3. 全局替换映射
    global_map = {}
    for k, v in config.get("全局替换", {}).items():
        global_map[str(k).strip()] = v

    # 4. 模板类型 → 模板根目录
    template_type = config.get("root", {}).get("模板类型", "")
    template_root = _resolve_template_root(params, template_type)
    if not template_root:
        stats = {
            "工程名称": project_name,
            "模板类型": template_type or "",
            "段落统计": {},
            "异常": (
                f"模板类型 '{template_type}' 未在全局配置中找到对应目录，"
                "且未提供 templates_root 参数"
            ),
            "堆栈": "",
        }
        write_error_reason(save_folder, stats, orig_name)
        return {"detail": "failed", "save_folder": save_folder, "stats": stats}

    # 5. 生成（捕获 MemoryError 与一般异常）
    try:
        stats = generate_pages_from_config(
            template_root=template_root,
            output_folder=save_folder,
            config=config,
            global_map=global_map,
            exempt_key=exempt_key,
        )
    except MemoryError:
        logger.exception("生成过程内存不足")
        stats = {
            "工程名称": project_name,
            "模板类型": template_type,
            "段落统计": {},
            "异常": "内存不足，请减少单次投递的数据量",
            "堆栈": traceback.format_exc(),
        }
        write_error_reason(save_folder, stats, orig_name)
        return {"detail": "failed", "save_folder": save_folder, "stats": stats}
    except Exception as e:
        logger.exception("生成过程异常")
        stats = {
            "工程名称": project_name,
            "模板类型": template_type,
            "段落统计": {},
            "异常": str(e),
            "堆栈": traceback.format_exc(),
        }
        write_error_reason(save_folder, stats, orig_name)
        return {"detail": "failed", "save_folder": save_folder, "stats": stats}

    stats["工程名称"] = project_name
    stats["模板类型"] = template_type
    detail = _judge_status(stats)

    # 6. 失败/部分成功补具体异常摘要（R2 + V-4）
    if detail in ("failed", "partial") and not stats.get("异常"):
        段落统计 = stats.get("段落统计", {})
        if not 段落统计:
            stats["异常"] = "配置文件中未识别到任何有效段落"
        else:
            fail_items = [
                (k, v.get("备注"), v.get("异常摘要", ""))
                for k, v in 段落统计.items()
                if v.get("备注") in _FAIL_NOTES
            ]
            if fail_items:
                summary = "；".join(
                    f"{k}：{note}" + (f"（{s}）" if s else "")
                    for k, note, s in fail_items
                )
                if detail == "failed":
                    stats["异常"] = f"所有段落均失败：{summary}"
                else:
                    stats["异常"] = f"部分段落失败：{summary}"
            elif detail == "failed":
                stats["异常"] = "所有段落数据条数为 0（无有效数据）"

    # 7. 三态产物写入
    if detail == "success":
        write_success_stats(save_folder, stats, orig_name)
        _safe_remove(os.path.join(save_folder, "错误原因.txt"))
    elif detail == "partial":
        write_success_stats(save_folder, stats, orig_name, partial=True)
        write_error_reason(save_folder, stats, orig_name, partial=True)
    else:
        write_error_reason(save_folder, stats, orig_name)
        _safe_remove(os.path.join(save_folder, "成功统计.txt"))

    # 8. 原配置档归档副本到工程目录（同名冲突加时间戳）
    try:
        _archive_conflict_free(txt_path, os.path.join(save_folder, orig_name))
    except OSError as e:
        logger.warning("配置档归档副本写入失败：%s", e)

    return {"detail": detail, "save_folder": save_folder, "stats": stats}


# ----------------------------------------------------------------------
# 插件入口
# ----------------------------------------------------------------------
def main(argv: list[str] | None = None) -> int:
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
        stream=sys.stderr,
    )

    argv = sys.argv if argv is None else argv
    if len(argv) < 2 or not argv[1].strip():
        print(
            json.dumps(
                {"status": "failed", "error": "缺少文件参数（argv[1]）"},
                ensure_ascii=False,
            )
        )
        return 0

    txt_path = argv[1]
    params = _load_params(argv)
    orig_name = os.path.basename(txt_path)

    try:
        outcome = process_config(txt_path, orig_name, params)
    except Exception as e:
        logger.exception("处理过程发生未预期异常")
        print(
            json.dumps(
                {
                    "status": "failed",
                    "plugin": "docgen",
                    "error": f"{type(e).__name__}: {e}",
                    "堆栈": traceback.format_exc(),
                },
                ensure_ascii=False,
            )
        )
        return 0

    detail = outcome["detail"]
    stats = outcome["stats"]
    段落统计 = stats.get("段落统计", {})
    generated = sum(1 for i in 段落统计.values() if i.get("数据条数", 0) > 0)

    result = {
        "status": "success" if detail in ("success", "partial") else "failed",
        "plugin": "docgen",
        "detail": detail,
        "partial": detail == "partial",
        "project": stats.get("工程名称", ""),
        "template_type": stats.get("模板类型", ""),
        "output_folder": outcome["save_folder"],
        "generated_sections": generated,
    }
    if detail == "failed":
        result["error"] = stats.get("异常", "未知异常")
        result["堆栈"] = stats.get("堆栈", "")

    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
